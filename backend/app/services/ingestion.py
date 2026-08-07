from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

TOKEN_RE = re.compile(r"[A-Za-z0-9]+|[\u4e00-\u9fff]")
SENTENCE_RE = re.compile(r"[^。！？!?；;.!?]+[。！？!?；;.!?]*")


@dataclass(slots=True)
class ParsedDocument:
    title: str
    text: str
    mime_type: str


@dataclass(slots=True)
class ChunkDraft:
    chunk_index: int
    title: str | None
    content: str
    token_count: int
    vector: dict[str, int]


def guess_mime_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    return {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".csv": "text/csv",
        ".json": "application/json",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pdf": "application/pdf",
    }.get(suffix, "application/octet-stream")


def parse_uploaded_file(filename: str, data: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = data.decode("utf-8", errors="ignore")
    elif suffix == ".csv":
        text = _parse_csv(data)
    elif suffix == ".json":
        text = _parse_json(data)
    elif suffix == ".docx":
        text = _parse_docx(data)
    elif suffix == ".pdf":
        text = _parse_pdf(data)
    else:
        text = data.decode("utf-8", errors="ignore")

    text = _normalize_text(text)
    return ParsedDocument(
        title=Path(filename).stem,
        text=text,
        mime_type=guess_mime_type(filename),
    )


def split_into_chunks(text: str, *, max_chars: int = 480, overlap: int = 80) -> list[ChunkDraft]:
    units = _segment_text(text, max_chars=max_chars)
    chunks: list[ChunkDraft] = []
    buffer: list[str] = []
    current_len = 0
    chunk_index = 0

    def flush() -> None:
        nonlocal buffer, current_len, chunk_index
        if not buffer:
            return
        content = "\n".join(buffer).strip()
        if not content:
            buffer = []
            current_len = 0
            return
        chunks.append(
            ChunkDraft(
                chunk_index=chunk_index,
                title=None,
                content=content,
                token_count=len(_tokenize(content)),
                vector=_vectorize(content),
            )
        )
        chunk_index += 1
        if overlap > 0:
            tail = content[-overlap:]
            buffer = [tail]
            current_len = len(tail)
        else:
            buffer = []
            current_len = 0

    for unit in units:
        unit_len = len(unit)
        if unit_len > max_chars:
            if buffer:
                flush()
            for piece in _split_long_text(unit, max_chars=max_chars, overlap=overlap):
                chunks.append(
                    ChunkDraft(
                        chunk_index=chunk_index,
                        title=None,
                        content=piece,
                        token_count=len(_tokenize(piece)),
                        vector=_vectorize(piece),
                    )
                )
                chunk_index += 1
            buffer = []
            current_len = 0
            continue

        if buffer and current_len + unit_len + 1 > max_chars:
            flush()
        buffer.append(unit)
        current_len += unit_len + 1

    flush()
    return chunks


def render_chunks_as_text(chunks: list[ChunkDraft]) -> str:
    parts: list[str] = []
    for chunk in chunks:
        parts.append(f"[Chunk {chunk.chunk_index}]\n{chunk.content}")
    return "\n\n".join(parts)


def _parse_csv(data: bytes) -> str:
    text_io = io.StringIO(data.decode("utf-8", errors="ignore"))
    reader = csv.reader(text_io)
    rows = list(reader)
    if not rows:
        return ""
    header = rows[0]
    body = rows[1:]
    lines = [", ".join(header)]
    for row in body:
        lines.append(", ".join(row))
    return "\n".join(lines)


def _parse_json(data: bytes) -> str:
    payload = json.loads(data.decode("utf-8", errors="ignore") or "{}")
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _parse_docx(data: bytes) -> str:
    document = DocxDocument(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _parse_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _normalize_text(text: str) -> str:
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _segment_text(text: str, *, max_chars: int) -> list[str]:
    segments: list[str] = []
    for paragraph in [paragraph.strip() for paragraph in re.split(r"\n\s*\n", text) if paragraph.strip()]:
        sentence_matches = [match.group().strip() for match in SENTENCE_RE.finditer(paragraph) if match.group().strip()]
        if not sentence_matches:
            sentence_matches = [paragraph]
        for sentence in sentence_matches:
            if len(sentence) <= max_chars:
                segments.append(sentence)
            else:
                segments.extend(_split_long_text(sentence, max_chars=max_chars, overlap=0))
    return segments


def _split_long_text(text: str, *, max_chars: int, overlap: int) -> list[str]:
    pieces: list[str] = []
    start = 0
    length = len(text)
    while start < length:
        end = min(length, start + max_chars)
        piece = text[start:end].strip()
        if piece:
            pieces.append(piece)
        if end >= length:
            break
        next_start = end - overlap if overlap > 0 else end
        if next_start <= start:
            next_start = end
        start = next_start
    return pieces


def _tokenize(text: str) -> list[str]:
    return [token.lower() for token in TOKEN_RE.findall(text)]


def _vectorize(text: str) -> dict[str, int]:
    vector: dict[str, int] = {}
    for token in _tokenize(text):
        vector[token] = vector.get(token, 0) + 1
    return vector
