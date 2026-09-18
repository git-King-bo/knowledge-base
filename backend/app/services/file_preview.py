import csv
from pathlib import Path

from openpyxl import load_workbook


def _table_page(rows, offset: int, limit: int):
    page = []
    total = 0
    columns = 0
    for number, row in enumerate(rows, 1):
        values = [str(value) if value is not None else "" for value in row]
        while values and not values[-1]:
            values.pop()
        if not any(value.strip() for value in values):
            continue
        columns = max(columns, len(values))
        if offset <= total < offset + limit:
            page.append({"number": number, "cells": values[:200]})
        total += 1
    return {"kind": "table", "rows": page, "total": total,
            "columns": min(columns, 200), "columns_truncated": columns > 200,
            "offset": offset, "limit": limit}


def preview_source(source, sheet: int, offset: int, limit: int):
    suffix = Path(source.filename).suffix.lower()
    if suffix == ".pdf":
        if not Path(source.storage_path).is_file():
            raise FileNotFoundError()
        return {"kind": "pdf"}
    if suffix == ".xlsx":
        workbook = load_workbook(source.storage_path, read_only=True, data_only=True)
        try:
            if sheet >= len(workbook.worksheets):
                raise ValueError("工作表不存在")
            result = _table_page(workbook.worksheets[sheet].iter_rows(values_only=True), offset, limit)
            return {**result, "sheets": workbook.sheetnames, "sheet": sheet}
        finally:
            workbook.close()
    if suffix == ".csv":
        with open(source.storage_path, encoding="utf-8-sig", errors="replace", newline="") as file:
            return {**_table_page(csv.reader(file), offset, limit), "sheets": [], "sheet": 0}
    text = source.content_text
    if not text:
        raise ValueError("暂无可预览的文本，请检查文件解析状态")
    # Text paging keeps long documents responsive without using retrieval chunks.
    start = offset * 20000
    return {"kind": "text", "text": text[start:start + 20000], "offset": offset,
            "total": max(1, (len(text) + 19999) // 20000),
            "note": "Word 展示提取的正文与表格文本，不保留原始排版。" if suffix == ".docx" else ""}
