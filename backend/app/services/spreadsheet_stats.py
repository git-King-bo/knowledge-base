"""Compute whole-workbook evidence for aggregate questions, outside chunk retrieval."""
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def needs_spreadsheet_stats(question: str) -> bool:
    return bool(re.search(
        r"多少|几[个人条行]|人数|数量|总数|总计|合计|一共|总共|统计|how\s+many|\bcount\b|\btotal\b",
        question, re.IGNORECASE,
    ))


def summarize_workbook(path: str) -> str:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        summaries = []
        for sheet in workbook.worksheets:
            header = None
            name_column = None
            row_count = 0
            name_count = 0
            names = set()
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if not any(values):
                    continue
                if header is None:
                    header = values
                    name_column = next((index for index, value in enumerate(values)
                                        if value.casefold() in {"姓名", "人员姓名", "员工姓名", "name", "full name"}), None)
                    continue
                row_count += 1
                if name_column is not None and name_column < len(values) and values[name_column]:
                    name_count += 1
                    names.add(values[name_column])
            summary = {
                "工作表": sheet.title,
                "非空行数（含首行）": row_count + int(header is not None),
                "首个非空行之后的非空行数": row_count,
            }
            if name_column is not None:
                summary.update({
                    "识别的表头": header,
                    "姓名非空的记录数（不含表头）": name_count,
                    "不同姓名文本数（不等于去重人数）": len(names),
                })
            summaries.append(summary)
        return json.dumps(summaries, ensure_ascii=False)
    finally:
        workbook.close()
