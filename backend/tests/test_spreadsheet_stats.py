import io
import json
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

import test_workspace as workspace
import test_streaming as streaming
from app.services.spreadsheet_stats import summarize_workbook


class SpreadsheetStatsTests(unittest.TestCase):
    setUp = workspace.WorkspaceTests.setUp
    tearDown = workspace.WorkspaceTests.tearDown
    base = workspace.WorkspaceTests.base
    fixture = streaming.StreamingTests.fixture

    def upload_roster(self, base, count=900, filename='名单.xlsx'):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '人员名单'
        sheet.append(['姓名', '个人简介'])
        for index in range(count):
            sheet.append([f'测试人员{index}', '第一行简介\n第二行简介'])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        response = self.client.post(f'/api/knowledge/bases/{base}/upload',
            files={'file': (filename, output.getvalue(), 'application/octet-stream')})
        self.assertEqual(response.status_code, 201, response.text)

    def test_total_uses_full_workbook_even_with_no_retrieved_chunks(self):
        base = self.base()
        self.upload_roster(base)
        other = self.base('其他库')
        self.upload_roster(other, 3, '不应出现.xlsx')
        with patch('app.api.routes.knowledge.KnowledgeIngestionRepository.search_chunks', return_value=[]):
            response = self.client.post('/api/knowledge/ask', json={
                'question': '一共涉及多少人员信息', 'knowledge_base_id': base, 'top_k': 1})
        self.assertEqual(response.status_code, 200, response.text)
        prompt = json.loads(self.requests[-1].content)['messages'][1]['content']
        self.assertIn('"姓名非空的记录数（不含表头）": 900', prompt)
        self.assertNotIn('不应出现.xlsx', prompt)
        self.client.put(f'/api/knowledge/bases/{other}', json={'status': 'archived'})
        response = self.client.post('/api/knowledge/ask', json={'question': '总人数'})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotIn('不应出现.xlsx', json.loads(self.requests[-1].content)['messages'][1]['content'])

    def test_streaming_receives_actual_count_not_filename_number(self):
        base = self.base()
        self.upload_roster(base, 111, '人才池-900.xlsx')
        upstream = 'data: {"choices":[{"delta":{"content":"111 条人员信息"}}]}\n\ndata: [DONE]\n\n'
        with self.fixture(upstream):
            response = self.client.post('/api/knowledge/ask/stream', json={
                'question': '一共射击多少人员信息', 'knowledge_base_id': base})
        self.assertEqual(response.status_code, 200, response.text)
        prompt = json.loads(self.requests[-1].content)['messages'][1]['content']
        self.assertIn('"姓名非空的记录数（不含表头）": 111', prompt)
        self.assertIn('event: done', response.text)

    def test_blank_rows_duplicate_names_and_multiple_sheets(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['姓名', '简介'])
        sheet.append(['甲', '多行\n内容'])
        sheet.append([None, None])
        sheet.append(['甲', '同名'])
        sheet.append([None, '未填写姓名'])
        sheet.cell(100, 1).number_format = '@'
        workbook.create_sheet('空表')
        second = workbook.create_sheet('产品表')
        second.append(['产品', '价格'])
        second.append(['设备', 0])
        path = Path(self.temp.name) / 'stats.xlsx'
        workbook.save(path)
        workbook.close()
        summaries = json.loads(summarize_workbook(str(path)))
        self.assertEqual(summaries[0]['姓名非空的记录数（不含表头）'], 2)
        self.assertEqual(summaries[0]['不同姓名文本数（不等于去重人数）'], 1)
        self.assertEqual(summaries[0]['首个非空行之后的非空行数'], 3)
        self.assertEqual(summaries[1]['非空行数（含首行）'], 0)
        self.assertNotIn('姓名非空的记录数（不含表头）', summaries[2])

    def test_unavailable_workbook_does_not_fabricate_total(self):
        base = self.base()
        self.upload_roster(base, 1)
        with patch('app.api.routes.knowledge.summarize_workbook', side_effect=OSError('missing')):
            response = self.client.post('/api/knowledge/ask', json={
                'question': '总人数', 'knowledge_base_id': base})
        self.assertEqual(response.status_code, 200, response.text)
        prompt = json.loads(self.requests[-1].content)['messages'][1]['content']
        self.assertIn('无法确认全表数量', prompt)
