"""Run with: python -m unittest discover -s tests -v (no paid services required)."""
import csv
import io
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from unittest.mock import patch

import httpx
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, inspect, text
from alembic import command
from alembic.config import Config
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.ai.providers.openai_compatible import OpenAICompatibleProvider
from app.core.config import settings
from app.db.models import AIActivityLogModel, AIProviderModel
from app.db.init_db import seed_db
from app.db.session import Base, get_db
from app.main import create_app
from app.repositories.sqlite import AIRepository
from app.schemas.usage import TokenUsage
from app.services.embeddings import EmbeddingClient


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.db.add(AIProviderModel(id='fixture', name='Fixture', provider='openai-compatible',
                    base_url='https://fixture.invalid/v1', default_model='fixture-model',
                    api_key_hint='test', is_default=True))
        self.db.add(AIProviderModel(id='mock', name='Mock', provider='mock',
                    base_url='local://mock', default_model='mock-chat', api_key_hint='', is_default=False))
        self.db.commit()
        app = create_app()
        def database():
            yield self.db
        app.dependency_overrides[get_db] = database
        # Startup is intentionally omitted: all tests use the isolated database above.
        self.client = TestClient(app)
        self.patches = [patch.object(settings, 'upload_dir', self.temp.name),
                        patch.object(settings, 'embedding_api_url', ''),
                        patch.object(settings, 'web_search_enabled', False)]
        for item in self.patches:
            item.start()
        self.requests = []
        def handler(request):
            self.requests.append(request)
            return httpx.Response(200, json={'choices': [{'message': {'content': '知识库测试答案'}}],
                'usage': {'prompt_tokens': 100, 'completion_tokens': 25, 'total_tokens': 125,
                          'prompt_tokens_details': {'cached_tokens': 40}}})
        self.transport = httpx.MockTransport(handler)
        self.adapter_patch = patch.object(OpenAICompatibleProvider, '_client',
            side_effect=lambda config: httpx.Client(base_url=config.base_url, transport=self.transport))
        self.adapter_patch.start()

    def tearDown(self):
        self.adapter_patch.stop()
        for item in reversed(self.patches):
            item.stop()
        self.client.close()
        self.db.close()
        self.engine.dispose()
        self.temp.cleanup()

    def base(self, name='研发知识库'):
        result = self.client.post('/api/knowledge/bases', json={'name': name, 'tags': ['测试']})
        self.assertEqual(result.status_code, 201, result.text)
        return result.json()['id']

    def upload(self, base, filename='流程.txt', content='产品发布流程：检查测试结果，审批通过后上线。'):
        return self.client.post(f'/api/knowledge/bases/{base}/upload',
                                files={'file': (filename, content.encode(), 'text/plain')})

    def test_end_to_end_knowledge_and_reported_usage(self):
        base = self.base()
        upload = self.upload(base)
        self.assertEqual(upload.status_code, 201, upload.text)
        source = upload.json()['source']
        self.assertGreater(source['chunk_count'], 0)
        search = self.client.get('/api/knowledge/search', params={'q': '产品发布流程', 'knowledge_base_id': base})
        self.assertEqual(search.status_code, 200, search.text)
        self.assertGreater(len(search.json()['hits']), 0)
        answer = self.client.post('/api/knowledge/ask', json={'question': '产品发布流程', 'knowledge_base_id': base})
        self.assertEqual(answer.status_code, 200, answer.text)
        self.assertEqual(answer.json()['answer'], '知识库测试答案')
        usage = self.client.get('/api/usage', params={'knowledge_base_id': base}).json()
        self.assertEqual(usage['summary']['total_tokens'], 125)
        self.assertEqual(usage['summary']['input_tokens'], 100)
        self.assertEqual(usage['summary']['output_tokens'], 25)
        self.assertEqual(usage['summary']['cached_tokens'], 40)
        self.assertEqual(usage['summary']['unknown_requests'], 0)
        self.assertEqual(usage['records'][0]['knowledge_base_id'], base)
        self.assertEqual(str(self.requests[0].url), 'https://fixture.invalid/v1/chat/completions')
        self.assertEqual(self.client.delete(f'/api/knowledge/bases/{base}/sources/{source["id"]}').status_code, 200)
        self.assertEqual(self.client.get('/api/knowledge/search', params={'q': '产品发布流程'}).json()['hits'], [])

    def test_old_unknown_mock_failed_and_connection_tests(self):
        repo = AIRepository(self.db)
        repo.create_activity_log(action='test', provider_id='fixture', model='fixture-model',
                                 request_text='', response_text='', success=True, latency_ms=1)
        # Legacy log has no usage row; migration must not invent a token count.
        self.db.add(AIActivityLogModel(id='legacy', action='chat', provider_id='fixture', model='old',
                    success=True, latency_ms=2, request_text='', response_text='', created_at=datetime.now(timezone.utc)))
        self.db.commit()
        response = self.client.post('/api/ai/chat', json={'provider_id': 'mock', 'messages': [{'role': 'user', 'content': 'test'}]})
        self.assertEqual(response.status_code, 200, response.text)
        with patch.object(OpenAICompatibleProvider, 'chat', side_effect=RuntimeError('fixture failure')):
            failed = self.client.post('/api/ai/chat', json={'messages': [{'role': 'user', 'content': 'test'}]})
        self.assertEqual(failed.status_code, 502)
        result = self.client.get('/api/usage').json()
        self.assertEqual(result['summary']['requests'], 3)
        self.assertEqual(result['summary']['unknown_requests'], 3)
        self.assertEqual(result['summary']['failures'], 1)
        self.assertEqual(result['summary']['total_tokens'], 0)
        self.assertEqual({item['source'] for item in result['records']}, {'unknown', 'mock'})
        self.assertTrue(all(item['total_tokens'] is None for item in result['records']))

    def test_xlsx_upload_indexes_all_nonempty_sheets(self):
        base = self.base()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '产品清单'
        sheet.append(['产品名称', '备注', '库存', '启用', '日期'])
        sheet.append(['星河设备', None, 0, False, datetime(2026, 9, 17)])
        workbook.create_sheet('空工作表')
        second = workbook.create_sheet('操作流程')
        second.append(['设备校准流程', '先检查电源再启动'])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        response = self.client.post(f'/api/knowledge/bases/{base}/upload',
            files={'file': ('设备.XLSX', output.getvalue(), 'application/octet-stream')})
        self.assertEqual(response.status_code, 201, response.text)
        source = response.json()['source']
        self.assertEqual(source['mime_type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertGreater(source['chunk_count'], 0)
        for query in ('星河设备', '设备校准流程'):
            hits = self.client.get('/api/knowledge/search',
                params={'q': query, 'knowledge_base_id': base}).json()['hits']
            self.assertTrue(hits, query)
        from app.services.ingestion import parse_uploaded_file
        parsed = parse_uploaded_file('设备.xlsx', output.getvalue()).text
        self.assertIn('工作表：产品清单', parsed)
        self.assertIn('工作表：操作流程', parsed)
        self.assertIn('星河设备 | | 0 | False | 2026-09-17', parsed)
        self.assertNotIn('空工作表', parsed)

    def test_empty_and_corrupt_xlsx_uploads_fail_cleanly(self):
        base = self.base()
        workbook = Workbook()
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        for content in (output.getvalue(), b'not an Excel workbook'):
            response = self.client.post(f'/api/knowledge/bases/{base}/upload',
                files={'file': ('empty.xlsx', content, 'application/octet-stream')})
            self.assertEqual(response.status_code, 400, response.text)
        sources = self.client.get(f'/api/knowledge/bases/{base}/sources').json()
        self.assertEqual(len(sources), 2)
        self.assertTrue(all(source['status'] == 'failed' for source in sources))

    def test_summary_is_not_limited_to_current_page_and_csv_matches(self):
        repo = AIRepository(self.db)
        for index in range(25):
            repo.create_activity_log(action='chat', provider_id='fixture', model='fixture-model',
                request_text='', response_text='', success=True, latency_ms=10,
                usage=TokenUsage(input_tokens=10, output_tokens=2, total_tokens=12, source='reported'))
        result = self.client.get('/api/usage', params={'page': 2, 'page_size': 20}).json()
        self.assertEqual(result['summary']['total_tokens'], 300)
        self.assertEqual(result['total'], 25)
        self.assertEqual(len(result['records']), 5)
        self.assertEqual(sum(point['total_tokens'] for point in result['daily']), 300)
        self.assertEqual(result['models'][0]['total_tokens'], 300)
        exported = self.client.get('/api/usage/export', params={'provider_id': 'fixture'})
        rows = list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
        self.assertEqual(len(rows), 25)
        self.assertEqual(sum(int(row['total_tokens']) for row in rows), 300)
        self.assertEqual(self.client.get('/api/usage', params={'model': 'missing'}).json()['total'], 0)
        self.assertEqual(self.client.get('/api/usage', params={'action': 'test'}).status_code, 422)

    def test_archiving_scope_and_file_validation(self):
        first, second = self.base('第一库'), self.base('第二库')
        self.assertEqual(self.upload(first).status_code, 201)
        self.assertEqual(self.client.get('/api/knowledge/search', params={'q': '发布', 'knowledge_base_id': second}).json()['hits'], [])
        self.assertEqual(self.upload(second, 'empty.txt', '').status_code, 400)
        self.assertEqual(self.upload(second, 'bad.exe', 'binary').status_code, 400)
        self.assertEqual(self.upload(second, 'bad.json', '{broken').status_code, 400)
        sources = self.client.get(f'/api/knowledge/bases/{second}/sources').json()
        self.assertEqual(sources[0]['status'], 'failed')
        self.client.put(f'/api/knowledge/bases/{first}', json={'status': 'archived'})
        self.assertEqual(self.upload(first).status_code, 409)
        self.assertEqual(self.client.get('/api/knowledge/search', params={'q': '发布', 'knowledge_base_id': first}).status_code, 409)
        self.assertEqual(self.client.get('/api/knowledge/search', params={'q': '发布'}).json()['hits'], [])
        self.assertEqual(self.client.get('/api/knowledge/search', params={'q': '发布', 'knowledge_base_id': 'missing'}).status_code, 404)
        self.assertEqual(self.client.get('/api/documents').status_code, 404)
        self.assertEqual(self.client.get('/api/categories').status_code, 404)

    def test_embedding_usage_is_recorded(self):
        base = self.base()
        def handler(request):
            import json
            size = len(json.loads(request.content)['input'])
            return httpx.Response(200, json={'data': [{'index': i, 'embedding': [1.0, 0.5]} for i in range(size)],
                                             'usage': {'prompt_tokens': 50, 'total_tokens': 50}})
        with patch.object(settings, 'embedding_api_url', 'https://embedding.invalid/v1'), patch.object(
            EmbeddingClient, '_client', side_effect=lambda: httpx.Client(base_url='https://embedding.invalid/v1',
                                                                         transport=httpx.MockTransport(handler))):
            uploaded = self.upload(base)
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        result = self.client.get('/api/usage', params={'action': 'embedding', 'knowledge_base_id': base}).json()
        self.assertEqual(result['summary']['total_tokens'], 50)
        self.assertEqual(result['summary']['input_tokens'], 50)
        self.assertEqual(result['summary']['output_tokens'], 0)
        self.assertEqual(result['records'][0]['source'], 'reported')

    def test_local_date_boundaries(self):
        local_today = (datetime.now(timezone.utc) + timedelta(hours=8)).date()
        boundary = datetime.combine(local_today, datetime.min.time()) - timedelta(hours=8)
        for identifier, created in [('before', boundary - timedelta(seconds=1)), ('inside', boundary + timedelta(seconds=1))]:
            self.db.add(AIActivityLogModel(id=identifier, action='chat', provider_id='fixture', model='test',
                                          success=True, latency_ms=0, request_text='', response_text='', created_at=created))
        self.db.commit()
        result = self.client.get('/api/usage', params={'days': 1, 'timezone_offset': 480}).json()
        self.assertEqual(result['total'], 1)
        self.assertEqual(result['records'][0]['id'], 'inside')
        self.assertEqual(result['daily'][0]['date'], local_today.isoformat())
        self.assertEqual(result['daily'][0]['requests'], 1)

    def test_bootstrap_preserves_saved_provider_and_default(self):
        with patch.object(settings, 'default_api_url', 'https://env.invalid/v1'):
            seed_db(self.db)
            provider = self.db.get(AIProviderModel, 'provider-agent-default')
            provider.default_model = 'saved-model'
            provider.base_url = 'https://saved.invalid/v1'
            self.db.commit()
            seed_db(self.db)
            self.assertEqual(provider.default_model, 'saved-model')
            self.assertEqual(provider.base_url, 'https://saved.invalid/v1')
            self.assertTrue(self.db.get(AIProviderModel, 'fixture').is_default)
            self.assertFalse(provider.is_default)

    def test_fresh_migration_and_repeat_upgrade(self):
        url = f'sqlite:///{self.temp.name}/migration.db'
        with patch.object(settings, 'database_url', url):
            command.upgrade(Config('alembic.ini'), 'head')
            command.upgrade(Config('alembic.ini'), 'head')
        engine = create_engine(url)
        self.assertTrue({'knowledge_bases', 'knowledge_sources', 'knowledge_chunks',
                         'ai_activity_logs', 'token_usage'}.issubset(set(inspect(engine).get_table_names())))
        engine.dispose()

    def test_migration_keeps_legacy_calls(self):
        url = f'sqlite:///{self.temp.name}/legacy.db'
        engine = create_engine(url)
        Base.metadata.create_all(engine)
        with engine.begin() as connection:
            connection.execute(text("DROP TABLE token_usage"))
            connection.execute(text("INSERT INTO ai_activity_logs "
                                    "(id, action, provider_id, model, success, latency_ms, request_text, response_text, created_at) "
                                    "VALUES ('legacy', 'chat', 'fixture', 'old', 1, 1, '', '', '2026-09-01 00:00:00')"))
        with patch.object(settings, 'database_url', url):
            command.stamp(Config('alembic.ini'), '0001_initial')
            command.upgrade(Config('alembic.ini'), 'head')
        with engine.connect() as connection:
            self.assertEqual(connection.scalar(text('SELECT count(*) FROM ai_activity_logs')), 1)
            self.assertEqual(connection.scalar(text('SELECT count(*) FROM token_usage')), 0)
        engine.dispose()

    def test_usage_parsing_does_not_estimate_missing_values(self):
        usage = TokenUsage.from_response({'usage': {'prompt_tokens': 100, 'completion_tokens': 20}})
        self.assertEqual(usage.total_tokens, 120)
        self.assertIsNone(usage.cached_tokens)
        self.assertIsNone(TokenUsage.from_response({}).total_tokens)
        self.assertIsNone(TokenUsage.from_response({'usage': {'prompt_tokens': -1}}).input_tokens)
        self.assertIsNone(TokenUsage.from_response({'usage': {'total_tokens': True}}).total_tokens)


if __name__ == '__main__':
    unittest.main()
